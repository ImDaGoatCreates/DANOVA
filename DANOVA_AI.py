from __future__ import annotations

import argparse
import json
import os
import platform
import queue
import re
import subprocess
import sys
import threading
import time
import urllib.parse
import webbrowser
import numpy as np
import sounddevice as sd
from dotenv import load_dotenv
from datetime import datetime
from pathlib import Path

import cv2
import geocoder
import requests
import speech_recognition as sr
from flask import Flask, jsonify, request

from PyQt6.QtCore import QThread, QTimer, Qt, pyqtSignal, QUrl
from PyQt6.QtGui import QImage, QPixmap, QFont
from PyQt6.QtWidgets import (
    QApplication,
    QCheckBox,
    QFileDialog,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)
from PyQt6.QtWebEngineWidgets import QWebEngineView

from robot_controller import HOME, RobotController
from danova_core import AuditLog, PiperSpeaker
from autodesk_cloud import DirectCloudClient

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
    from pptx import Presentation
except ImportError:
    pptx = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import whisper
except ImportError:
    whisper = None

stt_queue = queue.Queue()
load_dotenv()

ROOT = Path(__file__).resolve().parent
os.environ.setdefault("YOLO_CONFIG_DIR", str(ROOT / ".ultralytics"))
MEMORY_FILE = ROOT / "DANOVA_memory.json"
PROMPT_FILE = ROOT / "DANOVA_system_prompt.txt"
UPLOAD_DIR = ROOT / "uploads"
DEFAULT_LEADER_PORT = os.environ.get("DANOVA_LEADER_PORT", "COM12")
DEFAULT_FOLLOWER_PORT = os.environ.get("DANOVA_FOLLOWER_PORT", "COM11")
MODEL = os.environ.get("DANOVA_MODEL", "openai/gpt-oss-20b")
WHISPER_MODEL = os.environ.get("DANOVA_WHISPER_MODEL", "small")

DOWNLOADS_DIR = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Downloads"

ENDPOINTS = [
    os.environ.get("DANOVA_LLM_API_URL"),
    "http://127.0.0.1:1234/v1/chat/completions",
]
VISION_MODEL = ROOT / "yolov8n.pt"
YOLO26_DETECT_MODEL = ROOT / "yolo26n.pt"
YOLO26_POSE_MODEL = ROOT / "yolo26n-pose.pt"
HAND_MODEL = ROOT / "hand_landmarker.task"
SNAPSHOT_DIR = ROOT / "snapshots"
WAKE_WORDS_PATTERN = re.compile(
    r"\b(dan|danova|dan\s*ova|dan\s*nova|dan\s*over|dan\s*off\s*a|dan\s*of\s*a|than)\b",
    re.IGNORECASE
)

APP_ALIASES = {
    "chrome": "chrome",
    "google chrome": "chrome",
    "notepad": "notepad",
    "text editor": "notepad",
    "calculator": "calc",
    "calc": "calc",
    "file explorer": "explorer",
    "explorer": "explorer",
    "my computer": "explorer",
    "cmd": "cmd",
    "terminal": "cmd",
    "command prompt": "cmd",
    "edge": "msedge",
    "ms edge": "msedge",
    "browser": "msedge",
}


def analyze_file_content(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        return "File does not exist."

    ext = path.suffix.lower()

    try:
        if ext in [".txt", ".py", ".json", ".csv", ".md", ".cpp", ".js", ".html", ".css", ".c", ".h"]:
            return path.read_text(encoding="utf-8", errors="ignore")

        elif ext == ".pdf":
            if pypdf is None:
                return "PDF support requires 'pypdf'. Install it with: pip install pypdf"
            reader = pypdf.PdfReader(str(path))
            text_pages = []
            for idx, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text_pages.append(f"--- Page {idx + 1} ---\n{extracted}")
            return "\n\n".join(text_pages) if text_pages else "No extractable text found in PDF."

        elif ext == ".docx":
            if docx is None:
                return "Word document support requires 'python-docx'. Install it with: pip install python-docx"
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs) if paragraphs else "Word document is empty."

        elif ext == ".pptx":
            if pptx is None:
                return "PowerPoint support requires 'python-pptx'. Install it with: pip install python-pptx"
            prs = Presentation(str(path))
            slides_text = []
            for idx, slide in enumerate(prs.slides):
                slide_content = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text.strip())
                if slide_content:
                    slides_text.append(f"--- Slide {idx + 1} ---\n" + "\n".join(slide_content))
            return "\n\n".join(slides_text) if slides_text else "PowerPoint presentation is empty."

        else:
            return f"Unsupported file extension '{ext}' for text extraction."

    except Exception as e:
        return f"Error analyzing file content: {e}"


def create_word_document(title: str, content_blocks: list[dict], output_path: Path) -> str:
    if docx is None:
        return "Failed: 'python-docx' is not installed."

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)

    for block in content_blocks:
        b_type = block.get("type", "paragraph")
        text = block.get("text", "")

        if b_type == "heading1":
            doc.add_heading(text, level=1)
        elif b_type == "heading2":
            doc.add_heading(text, level=2)
        elif b_type == "bullet":
            doc.add_paragraph(text, style='List Bullet')
        else:
            doc.add_paragraph(text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return f"Word document created successfully at {output_path}"


def create_powerpoint_presentation(title: str, slides_data: list[dict], output_path: Path) -> str:
    if pptx is None:
        return "Failed: 'python-pptx' is not installed."

    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title

    bullet_slide_layout = prs.slide_layouts[1]
    for slide_info in slides_data:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]

        title_shape.text = slide_info.get("title", "Slide")
        tf = body_shape.text_frame

        bullets = slide_info.get("bullets", [])
        if bullets:
            tf.text = bullets[0]
            for bullet in bullets[1:]:
                p = tf.add_paragraph()
                p.text = bullet

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return f"PowerPoint presentation created successfully at {output_path}"


def load_memory() -> list[dict]:
    try:
        raw = json.loads(MEMORY_FILE.read_text(encoding="utf-8"))
        return [
            x if isinstance(x, dict) else {"role": "assistant", "content": str(x)}
            for x in raw
        ]
    except (OSError, json.JSONDecodeError):
        return []

def sanitize_voice_text(text: str) -> str:
    text = re.sub(r'```.*?```', '', text, flags=re.DOTALL)
    text = re.sub(r"[^a-zA-Z0-9\s.,':;()]", "", text)
    return re.sub(r"\s+", " ", text).strip()

def extract_wake_word_command(text: str) -> str | None:
    match = WAKE_WORDS_PATTERN.search(text)
    if not match:
        return None
    command = text[match.end():].strip()
    return command if command else ""

def clean_speech(text: str) -> str:
    text = text.lower().strip()

    text = re.sub(r"\b(hey|dan|okay|please)\b", "", text)

    text = text.replace("pull up", "open")
    text = text.replace("bring up", "open")

    text = text.replace(" on ", " in ")
    text = text.replace(" for ", " in ")

    text = re.sub(r"\b(my|the|a|an|design|designs)\b", "", text)

    text = re.sub(r"[^\w\s\-']", "", text)

    text = re.sub(r"\s+", " ", text)

    return text.strip()


def parse_cloud_command(text: str):
    """Parse a voice command for cloud design requests.

    Returns a tuple (folder_query, file_name).  ``file_name`` is empty when the
    request refers to a folder only.  The function normalises common phrases so
    that callers can rely on a consistent format.
    """
    text = text.lower()

    for old, new in ("pull up", "bring up"), ("open", "bring up"), ("show", "bring up"):
        text = text.replace(old, new)

    file_match = re.search(r"bring up\s+(.+?)\s+(?:from|in)\s+(.*)", text)
    if file_match:
        file_name = file_match.group(1).strip()
        folder_name = file_match.group(2).strip()
        return folder_name, file_name

    folder_match = re.search(r"bring up\s+(?:my\s+)?designs?\s+(?:for|in)\s+(.*)", text)
    if folder_match:
        folder_name = folder_match.group(1).strip()
        return folder_name, ""

    return text.strip(), ""

def save_turn(user: str, assistant: str) -> None:
    memory = load_memory()[-78:]
    stamp = datetime.now().isoformat(timespec="seconds")
    memory += [
        {"role": "user", "content": user, "time": stamp},
        {"role": "assistant", "content": assistant, "time": stamp},
    ]
    temp = MEMORY_FILE.with_suffix(".tmp")
    temp.write_text(json.dumps(memory, indent=2, ensure_ascii=False), encoding="utf-8")
    temp.replace(MEMORY_FILE)


def system_prompt() -> str:
    base_prompt = ""
    try:
        base_prompt = PROMPT_FILE.read_text(encoding="utf-8")
    except OSError:
        base_prompt = "You are D.A.N.O.V.A., a helpful voice desktop assistant."

    doc_instructions = (
        "\n\n[Document & Presentation Generation Instructions]:\n"
        "If requested to create a PowerPoint presentation, format your response with a code block:\n"
        "```json_pptx\n"
        "{\n"
        '  "title": "Presentation Title",\n'
        '  "slides": [\n'
        '    {"title": "Slide Title", "bullets": ["Bullet 1", "Bullet 2"]}\n'
        "  ]\n"
        "}\n"
        "```\n"
        "If requested to create a Word document, format your response with a code block:\n"
        "```json_docx\n"
        "{\n"
        '  "title": "Document Title",\n'
        '  "blocks": [\n'
        '    {"type": "heading1", "text": "Heading text"},\n'
        '    {"type": "paragraph", "text": "Body paragraph text"},\n'
        '    {"type": "bullet", "text": "Bullet point text"}\n'
        "  ]\n"
        "}\n"
        "```\n"
    )
    return base_prompt + doc_instructions


def endpoint() -> str | None:
    for url in dict.fromkeys(x.rstrip("/") for x in ENDPOINTS if x):
        try:
            if requests.get(url.rsplit("/v1/", 1)[0] + "/v1/models", timeout=2).ok:
                return url
        except requests.RequestException:
            pass
    return None


def open_app(app_name: str) -> str:
    clean_name = app_name.lower().strip()
    target = APP_ALIASES.get(clean_name, clean_name)

    try:
        if platform.system() == "Windows":
            os.system(f"start {target}")
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", "-a", target])
        else:
            subprocess.Popen([target])
        return f"Opening {app_name}."
    except Exception as e:
        return f"Failed to launch {app_name}. Error: {e}"


def extract_code_blocks(text: str) -> list[tuple[str, str]]:
    pattern = re.compile(r"```(?P<lang>[\w+-]+)?\n(?P<code>.*?)```", re.DOTALL)
    blocks = []
    for match in pattern.finditer(text):
        lang = match.group("lang") or "txt"
        code = match.group("code")
        ext_map = {"python": "py", "py": "py", "cpp": "cpp", "c": "c", "json": "json", "html": "html", "javascript": "js", "js": "js"}
        ext = ext_map.get(lang.lower(), "txt")
        filename = f"generated_output_{int(time.time())}.{ext}"
        blocks.append((filename, code))
    return blocks


class STTThread(threading.Thread):
    def __init__(self):
        super().__init__()
        self.running = True

        if whisper is None:
            raise RuntimeError(
                "Whisper is not installed or could not be imported. Install it with `pip install openai-whisper` "
                "and ensure the environment uses the same Python interpreter as this app."
            )

        model_name = WHISPER_MODEL
        try:
            self.model = whisper.load_model(model_name)
        except Exception as e:
            print(f"[WHISPER] Failed to load model '{model_name}': {e}")
            print("[WHISPER] Falling back to 'tiny'.")
            self.model = whisper.load_model("tiny")

        self.sample_rate = 16000
        self.buffer = []
        self.max_buffer_seconds = 6

    def run(self):
        chunk_size = int(0.25 * self.sample_rate)

        silence_threshold = 0.01
        silence_time_required = 1.2
        min_record_time = 1.8

        silence_start = None
        speech_detected = False
        recording_start = None

        stream = sd.InputStream(
            samplerate=self.sample_rate,
            channels=1,
            dtype="float32",
            blocksize=chunk_size
        )
        stream.start()

        while self.running:
            try:
                chunk, _ = stream.read(chunk_size)
                volume = float(np.abs(chunk).mean())

                self.buffer.append(chunk)
                max_chunks = int(10 / 0.25)
                if len(self.buffer) > max_chunks:
                    self.buffer = self.buffer[-max_chunks:]

                if volume > silence_threshold:
                    silence_start = None

                    if not speech_detected:
                        speech_detected = True
                        recording_start = time.time()

                        try:
                            if 'window' in globals():
                                window.status_label.setText("Listening")
                        except:
                            pass

                else:
                    if silence_start is None:
                        silence_start = time.time()

                    elif time.time() - silence_start > silence_time_required:

                        if not speech_detected:
                            self.buffer = []
                            continue

                        if recording_start and (time.time() - recording_start < min_record_time):
                            continue

                        audio = np.concatenate(self.buffer)
                        self.buffer = []
                        speech_detected = False

                        if len(audio) < self.sample_rate * 0.6:
                            continue

                        max_samples = self.sample_rate * 10
                        if len(audio) > max_samples:
                            audio = audio[-max_samples:]

                        audio = audio.flatten().astype("float32")

                        max_val = max(1e-6, float(np.max(np.abs(audio))))
                        audio = audio / max_val

                        rms = float(np.sqrt(np.mean(audio**2)))
                        if rms < 0.015:
                            continue

                        try:
                            if 'window' in globals():
                                window.status_label.setText("Thinking")
                        except:
                            pass

                        try:
                            result = self.model.transcribe(
                                audio,
                                verbose=False,
                                language="en",
                                condition_on_previous_text=False,
                                temperature=0.0,
                                beam_size=8,
                                best_of=8
                            )
                            text = result.get("text", "").strip()
                        except Exception:
                            continue

                        print("[HEARD]", text)

                        if not text:
                            continue

                        text_check = re.sub(r"[^\w\s']", "", text.lower())

                        if text_check in ["you", "boom"]:
                            continue

                        if not any(c.isalpha() for c in text_check):
                            continue

                        stt_queue.put(text)

                        silence_start = None

            except Exception as e:
                print("[STT ERROR]", e)
                continue

        stream.stop()

    def stop(self):
        self.running = False


class DanWindow(QMainWindow):
    update_editor_signal = pyqtSignal(str)
    update_viewer_signal = pyqtSignal(str)
    update_label_signal = pyqtSignal(str)
    update_file_label_signal = pyqtSignal(str)
    update_path_label_signal = pyqtSignal(str)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)

        self.robot = RobotController(DEFAULT_LEADER_PORT, DEFAULT_FOLLOWER_PORT)
        self.speaker = PiperSpeaker()
        self.cloud_client = DirectCloudClient()
        self.vision_context = "Camera offline"
        self.person_detected = False
        self.cloud_client.login()
        self.cloud_files = []

        self.location: dict | None = None
        self.cached_weather: str = "Weather data unavailable."
        self.file_context: str = ""
        self.active_file_path: str = ""
        self.last_ai_response: str = ""

        self.setWindowTitle("D.A.N.O.V.A. | Digital Autonomous Neural Operations & Virtual Assistant")
        self.resize(1300, 800)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        self._build_ui(central_widget)

        self.update_editor_signal.connect(self.code_editor.setPlainText)
        self.update_viewer_signal.connect(self.load_interactive_cad_viewer)
        self.update_label_signal.connect(self.viewer_info_label.setText)
        self.update_file_label_signal.connect(self.file_label.setText)
        self.update_path_label_signal.connect(self.editor_path_label.setText)

        self.refresh_status()
        self.status_timer = QTimer(self)
        self.status_timer.timeout.connect(self.refresh_status)
        self.status_timer.start(2500)

        self._start_gps_updates()
        self._start_upload_server()

    def _build_ui(self, central_widget):
        root = QVBoxLayout(central_widget)

        title = QLabel("D.A.N.O.V.A. | Digital Autonomous Neural Operations & Virtual Assistant")
        title.setStyleSheet("font-size:17px; font-weight:bold; color:#51d8ff;")
        root.addWidget(title)

        top_bar = QHBoxLayout()
        self.status_label = QLabel("Listening")
        self.status_label.setStyleSheet("font-size:14px; color:#00ffcc; font-weight:bold;")
        top_bar.addWidget(self.status_label)
        root.addLayout(top_bar)

        self.code_editor = QTextEdit()

        content = QHBoxLayout()

        left_col = QWidget()
        left_layout = QVBoxLayout(left_col)
        left_layout.addWidget(QLabel("VISION // LOCAL CAMERA"))
        self.camera_view = QLabel("Camera is off")
        self.camera_view.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.camera_view.setMinimumSize(320, 200)
        self.camera_view.setStyleSheet(
            "background:#080b10; color:#7a9aae; border:1px solid #294355;"
        )
        left_layout.addWidget(self.camera_view)

        self.vision_status = QLabel("Camera offline")
        self.vision_status.setWordWrap(True)
        left_layout.addWidget(self.vision_status)

        vision_row = QHBoxLayout()
        vision_row.addWidget(QLabel("Camera"))
        self.camera_index = QLineEdit("0")
        self.camera_index.setMaximumWidth(45)
        vision_row.addWidget(self.camera_index)
        self.detector_box = QCheckBox("Object detection")
        self.detector_box.setChecked(True)
        vision_row.addWidget(self.detector_box)
        left_layout.addLayout(vision_row)

        self.start_camera_btn = QPushButton("Start camera")
        self.start_camera_btn.clicked.connect(self.start_camera)
        self.stop_camera_btn = QPushButton("Stop camera")
        self.stop_camera_btn.clicked.connect(self.stop_camera)

        left_layout.addWidget(self.start_camera_btn)
        left_layout.addWidget(self.stop_camera_btn)
        left_layout.addSpacing(10)

        left_layout.addWidget(QLabel("ROBOTICS & FILE WORKSPACE"))
        file_box = QHBoxLayout()
        self.upload_btn = QPushButton("Upload File")
        self.upload_btn.clicked.connect(self.select_file)
        self.save_btn = QPushButton("Save to Downloads")
        self.save_btn.clicked.connect(self.save_extracted_file)
        file_box.addWidget(self.upload_btn)
        file_box.addWidget(self.save_btn)
        left_layout.addLayout(file_box)

        self.file_label = QLabel("No file loaded.")
        self.file_label.setStyleSheet("color: #7a9aae; font-size: 11px;")
        left_layout.addWidget(self.file_label)

        self.robot_status = QLabel()
        self.robot_status.setWordWrap(True)
        left_layout.addWidget(self.robot_status)

        ports = QGridLayout()
        ports.addWidget(QLabel("Leader"), 0, 0)
        self.leader_port = QLineEdit(DEFAULT_LEADER_PORT)
        ports.addWidget(self.leader_port, 0, 1)
        ports.addWidget(QLabel("Follower"), 1, 0)
        self.follower_port = QLineEdit(DEFAULT_FOLLOWER_PORT)
        ports.addWidget(self.follower_port, 1, 1)
        left_layout.addLayout(ports)

        for label, method in [
            ("Connect & verify", self.connect_arms),
            ("Safe home", self.home_arms),
            ("Torque OFF", lambda: self.set_torque(False)),
            ("Torque ON", lambda: self.set_torque(True)),
        ]:
            button = QPushButton(label)
            button.clicked.connect(method)
            left_layout.addWidget(button)

        left_layout.addStretch()
        content.addWidget(left_col, 2)

        middle_col = QWidget()
        middle_layout = QVBoxLayout(middle_col)
        
        editor_header_layout = QHBoxLayout()
        editor_title = QLabel("BUILT-IN PROGRAM & TEXT EDITOR")
        editor_title.setStyleSheet("color: #51d8ff; font-weight: bold;")
        editor_header_layout.addWidget(editor_title)
        
        self.save_editor_btn = QPushButton("Save File Changes")
        self.save_editor_btn.clicked.connect(self.save_editor_content)
        self.save_editor_btn.setStyleSheet("background: #00ffcc; color: #000; font-weight: bold; padding: 4px 10px;")
        editor_header_layout.addWidget(self.save_editor_btn)
        middle_layout.addLayout(editor_header_layout)

        self.editor_path_label = QLabel("Active File: None")
        self.editor_path_label.setStyleSheet("color: #7a9aae; font-size: 11px;")
        middle_layout.addWidget(self.editor_path_label)

        self.code_editor = QTextEdit()
        self.code_editor.setFont(QFont("Consolas", 11))
        self.code_editor.setStyleSheet(
            "background: #0d131d; color: #00ffcc; border: 1px solid #294355; selection-background-color: #294355;"
        )
        self.code_editor.setPlaceholderText("Select a file using workspace explorer or voice command to view and edit code directly here...")
        middle_layout.addWidget(self.code_editor)

        content.addWidget(middle_col, 3)

        right_col = QWidget()
        right_layout = QVBoxLayout(right_col)

        viewer_header_layout = QHBoxLayout()
        viewer_title = QLabel("FUSION 360 CLOUD 3D VIEWER")
        viewer_title.setStyleSheet("color: #51d8ff; font-weight: bold;")
        viewer_header_layout.addWidget(viewer_title)

        self.load_cloud_cad_btn = QPushButton("Load Fusion Cloud Design")
        self.load_cloud_cad_btn.clicked.connect(self.prompt_cloud_cad_url)
        self.load_cloud_cad_btn.setStyleSheet("background: #51d8ff; color: #000; font-weight: bold; padding: 4px 10px;")
        viewer_header_layout.addWidget(self.load_cloud_cad_btn)
        right_layout.addLayout(viewer_header_layout)

        self.viewer_info_label = QLabel("Navigation: Click & Drag to Rotate | Scroll to Zoom | Right-Click to Pan")
        self.viewer_info_label.setStyleSheet("color: #7a9aae; font-size: 11px;")
        right_layout.addWidget(self.viewer_info_label)

        self.cloud_cad_view = QWebEngineView()
        self.load_interactive_cad_viewer("")
        right_layout.addWidget(self.cloud_cad_view)

        content.addWidget(right_col, 4)
        root.addLayout(content)

    def load_interactive_cad_viewer(self, urn: str = ""):
        if urn:
            token = self.cloud_client.access_token or ""
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta name="viewport" content="width=device-width, initial-scale=1.0">
                <link rel="stylesheet" href="https://developer.api.autodesk.com/modelderivative/v2/viewers/7.*/style.min.css" type="text/css">
                <script src="https://developer.api.autodesk.com/modelderivative/v2/viewers/7.*/viewer3D.min.js"></script>
                <style> body, html {{ margin: 0; height: 100%; overflow: hidden; background: #080b10; }} #viewerDiv {{ width: 100%; height: 100%; }} </style>
            </head>
            <body>
                <div id="viewerDiv"></div>
                <script>
                    var viewer;
                    var options = {{
                        env: 'AutodeskProduction',
                        getAccessToken: function(onTokenReady) {{
                            onTokenReady('{token}', 3600);
                        }}
                    }};

                    Autodesk.Viewing.Initializer(options, function() {{
                        var htmlDiv = document.getElementById('viewerDiv');
                        viewer = new Autodesk.Viewing.GuiViewer3D(htmlDiv);
                        viewer.start();
                        var documentId = 'urn:' + '{urn}';
                        Autodesk.Viewing.Document.load(documentId, onDocumentLoadSuccess, onDocumentLoadFailure);
                    }});

                    function onDocumentLoadSuccess(doc) {{
                        var defaultModel = doc.getRoot().getDefaultGeometry();
                        viewer.loadDocumentNode(doc, defaultModel);
                    }}

                    function onDocumentLoadFailure() {{
                        console.error('Failed fetching CAD model from Autodesk Cloud.');
                    }}
                </script>
            </body>
            </html>
            """
            self.cloud_cad_view.setHtml(html_content, QUrl("http://localhost"))
        else:
            html_content = """
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>Fusion 360 Cloud 3D Viewer</title>
                <style>
                    body, html { margin: 0; padding: 0; width: 100%; height: 100%; overflow: hidden; background: #080b10; font-family: monospace; color: #7a9aae; }
                </style>
                <script src="https://cdnjs.cloudflare.com/ajax/libs/three.js/r128/three.min.js"></script>
                <script src="https://cdn.jsdelivr.net/npm/three@0.128.0/examples/js/controls/OrbitControls.js"></script>
            </head>
            <body>
                <div id="canvas-container"></div>
                <script>
                    const container = document.getElementById('canvas-container');
                    const scene = new THREE.Scene();
                    scene.background = new THREE.Color(0x080b10);

                    const camera = new THREE.PerspectiveCamera(45, window.innerWidth / window.innerHeight, 0.1, 1000);
                    camera.position.set(25, 20, 25);

                    const renderer = new THREE.WebGLRenderer({ antialias: true });
                    renderer.setSize(window.innerWidth, window.innerHeight);
                    container.appendChild(renderer.domElement);

                    const controls = new THREE.OrbitControls(camera, renderer.domElement);
                    controls.enableDamping = true;
                    controls.dampingFactor = 0.05;

                    const ambientLight = new THREE.AmbientLight(0xffffff, 0.7);
                    scene.add(ambientLight);

                    const dirLight = new THREE.DirectionalLight(0xffffff, 0.8);
                    dirLight.position.set(20, 40, 20);
                    scene.add(dirLight);

                    const group = new THREE.Group();
                    const baseGeo = new THREE.BoxGeometry(12, 1, 8);
                    const mat = new THREE.MeshStandardMaterial({ color: 0x1d3557, metalness: 0.5, roughness: 0.3 });
                    const base = new THREE.Mesh(baseGeo, mat);
                    base.position.y = -0.5;
                    group.add(base);

                    const cylGeo = new THREE.CylinderGeometry(2, 2, 6, 32);
                    const cylMat = new THREE.MeshStandardMaterial({ color: 0x51d8ff, metalness: 0.8, roughness: 0.2 });
                    const cyl = new THREE.Mesh(cylGeo, cylMat);
                    cyl.position.set(0, 3, 0);
                    group.add(cyl);

                    scene.add(group);

                    window.addEventListener('resize', () => {
                        camera.aspect = window.innerWidth / window.innerHeight;
                        camera.updateProjectionMatrix();
                        renderer.setSize(window.innerWidth, window.innerHeight);
                    });

                    function animate() {
                        requestAnimationFrame(animate);
                        controls.update();
                        renderer.render(scene, camera);
                    }
                    animate();
                </script>
            </body>
            </html>
            """
            self.cloud_cad_view.setHtml(html_content, QUrl("http://localhost"))

    def handle_file_double_click(self, event):
        cursor = self.code_editor.cursorForPosition(event.pos())
        cursor.select(cursor.SelectionType.LineUnderCursor)
        line = cursor.selectedText().strip()

        match = re.match(r"\[(\d+)\]\s+(.+)", line)
        if not match:
            return

        index = int(match.group(1))

        if not hasattr(self, "cloud_files"):
            return

        if index < 0 or index >= len(self.cloud_files):
            return

        file = self.cloud_files[index]

        try:
            urn = self.cloud_client.find_design_urn_globally(file["name"])
            if urn:
                self.load_interactive_cad_viewer(urn)
                self.viewer_info_label.setText(f"Loaded: {file['name']}")
            else:
                QMessageBox.warning(self, "Error", "Could not load selected design.")
        except Exception as e:
            QMessageBox.warning(self, "Error", str(e))

    def prompt_cloud_cad_url(self):
        from PyQt6.QtWidgets import QInputDialog
        design_name, ok = QInputDialog.getText(self, "Load Fusion Cloud Design", "Enter Design Name:")
        if ok and design_name.strip():
            try:
                urn = self.cloud_client.find_design_urn_globally(design_name.strip())
                if urn:
                    self.load_interactive_cad_viewer(urn)
                    self.viewer_info_label.setText(f"Loaded Cloud Design: {design_name.strip()} (URN: {urn})")
                else:
                    QMessageBox.critical(self, "Not Found", f"Could not find design '{design_name.strip()}' in your cloud projects.")
            except Exception as e:
                QMessageBox.critical(self, "Cloud Error", str(e))

    def select_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Select File for Context", "", "Supported Files (*.txt *.pdf *.docx *.pptx *.py *.json *.csv *.md *.cpp *.js *.html *.css);;All Files (*)"
        )
        if file_path:
            self.load_and_process_file(file_path)

    def load_and_process_file(self, file_path: str):
        try:
            content = analyze_file_content(file_path)
            self.active_file_path = file_path
            self.file_context = f"\n\n[File Contents of {Path(file_path).name}]:\n{content}"
            self.file_label.setText(f"Loaded: {Path(file_path).name}")
            self.editor_path_label.setText(f"Active File: {file_path}")
            self.code_editor.setPlainText(content)
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to read file: {e}")

    def save_editor_content(self):
        if not self.active_file_path:
            file_path, _ = QFileDialog.getSaveFileName(
                self, "Save Code As", str(DOWNLOADS_DIR), "Text Files (*.txt *.py *.json *.csv *.md *.cpp *.js);;All Files (*)"
            )
            if not file_path:
                return
            self.active_file_path = file_path
            self.editor_path_label.setText(f"Active File: {file_path}")

        try:
            content = self.code_editor.toPlainText()
            Path(self.active_file_path).write_text(content, encoding="utf-8")
            QMessageBox.information(self, "Success", f"Successfully saved file to:\n{self.active_file_path}")
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to save file: {e}")

    def save_extracted_file(self):
        if not self.last_ai_response:
            QMessageBox.information(self, "Info", "No recent AI response to save.")
            return

        pptx_match = re.search(r"```json_pptx\n(.*?)\n```", self.last_ai_response, re.DOTALL)
        if pptx_match:
            try:
                data = json.loads(pptx_match.group(1))
                dest_path = DOWNLOADS_DIR / f"presentation_{int(time.time())}.pptx"
                msg = create_powerpoint_presentation(data.get("title", "Presentation"), data.get("slides", []), dest_path)
                QMessageBox.information(self, "Success", msg)
                return
            except Exception as e:
                QMessageBox.warning(self, "Error", f"Failed to generate PowerPoint: {e}")

        docx_match = re.search(r"```json_docx\n(.*?)\n```", self.last_ai_response, re.DOTALL)
        if docx_match:
            try:
                data = json.loads(docx_match.group(1))
                dest_path = DOWNLOADS_DIR / f"document_{int(time.time())}.docx"
                msg = create_word_document(data.get("title", "Document"), data.get("blocks", []), dest_path)
                QMessageBox.information(self, "Success", msg)
                return
            except Exception as e:
                QMessageBox.warning(self, "Error", f"Failed to generate Word document: {e}")

        blocks = extract_code_blocks(self.last_ai_response)
        if not blocks:
            QMessageBox.information(self, "Info", "No code blocks or generated documents found in latest AI output.")
            return

        DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
        saved_paths = []
        for filename, content in blocks:
            dest_path = DOWNLOADS_DIR / filename
            dest_path.write_text(content, encoding="utf-8")
            saved_paths.append(str(dest_path))

        if saved_paths:
            self.active_file_path = saved_paths[0]
            self.editor_path_label.setText(f"Active File: {self.active_file_path}")
            self.code_editor.setPlainText(blocks[0][1])

        QMessageBox.information(self, "Success", f"Saved output file(s) and loaded into built-in editor:\n" + "\n".join(saved_paths))

    def get_context_info(self) -> str:
        now = datetime.now()
        date_str = now.strftime("%A, %B %d, %Y")
        time_str = now.strftime("%I:%M %p")

        loc_str = "Unknown location"
        if self.location:
            loc_str = f"Lat: {self.location.get('lat')}, Lon: {self.location.get('lon')}"

        return (
            f"[Current Date]: {date_str}\n"
            f"[Current Time]: {time_str}\n"
            f"[GPS Location]: {loc_str}\n"
            f"[Current Weather]: {self.cached_weather}"
            f"{getattr(self, 'file_context', '')}"
        )

    def fetch_weather(self):
        if not self.location:
            return
        try:
            lat = self.location.get("lat")
            lon = self.location.get("lon")
            url = f"https://api.open-meteo.com/v1/forecast?latitude={lat}&longitude={lon}&current_weather=true"
            res = requests.get(url, timeout=5)
            if res.ok:
                data = res.json().get("current_weather", {})
                temp_c = data.get("temperature")
                temp_f = (temp_c * 9 / 5) + 32 if temp_c is not None else "N/A"
                wind = data.get("windspeed")
                self.cached_weather = f"{temp_f:.1f}°F ({temp_c}°C), Wind Speed: {wind} km/h"
        except Exception as e:
            print(f"[Weather Error] {e}")

    def generate_wake_greeting(self) -> str:
        greeting_prompt = (
            "[System Event]: The user just spoke your wake word to get your attention. "
            "Acknowledge them with a short, natural, single-sentence response (e.g., 'Yes?', 'How can I help you?')."
        )
        return self.ask_assistant(greeting_prompt)

    def ask_assistant(self, user_text: str) -> str:
        messages = [{"role": "system", "content": system_prompt()}]
        messages += [
            {"role": x.get("role", "assistant"), "content": x.get("content", "")}
            for x in load_memory()[-16:]
            if x.get("content")
        ]

        environment_context = self.get_context_info()
        full_context = f"\n\n[Environment Context]:\n{environment_context}"
        if self.vision_context:
            full_context += f"\n[Camera Vision]: {self.vision_context}"

        messages.append({"role": "user", "content": user_text + full_context})

        url = endpoint()
        if not url:
            return "I cannot reach the local language model."

        try:
            response = requests.post(
                url,
                json={"model": MODEL, "messages": messages, "temperature": 0.6},
                timeout=500,
            )
            response.raise_for_status()
            data = response.json()

            if "choices" in data:
                text = data["choices"][0]["message"]["content"].strip()
            elif "message" in data and "content" in data["message"]:
                text = data["message"]["content"].strip()
            elif "response" in data:
                text = data["response"].strip()
            else:
                text = "I received an unexpected response format."
        except Exception:
            text = "The language model returned an error."

        self.last_ai_response = text

        if "```json_pptx" in text:
            pptx_match = re.search(r"```json_pptx\n(.*?)\n```", text, re.DOTALL)
            if pptx_match:
                try:
                    p_data = json.loads(pptx_match.group(1))
                    dest_file = DOWNLOADS_DIR / f"presentation_{int(time.time())}.pptx"
                    create_powerpoint_presentation(p_data.get("title", "Presentation"), p_data.get("slides", []), dest_file)
                    text += f"\n\n[Created PowerPoint presentation at {dest_file.name}]"
                except Exception as e:
                    print(f"[PPTX Error]: {e}")

        elif "```json_docx" in text:
            docx_match = re.search(r"```json_docx\n(.*?)\n```", text, re.DOTALL)
            if docx_match:
                try:
                    d_data = json.loads(docx_match.group(1))
                    dest_file = DOWNLOADS_DIR / f"document_{int(time.time())}.docx"
                    create_word_document(d_data.get("title", "Document"), d_data.get("blocks", []), dest_file)
                    text += f"\n\n[Created Word document at {dest_file.name}]"
                except Exception as e:
                    print(f"[DOCX Error]: {e}")

        save_turn(user_text, text)
        return text

    def play_media(self, query: str) -> str:
        if not query:
            return "Please specify what song, video, or topic you want to play."

        if "drop my needle" in query.lower() or "drop the needle" in query.lower():
            current_month = datetime.now().month
            if current_month == 12:
                query = "Jingle Bells Bombay Dub Orchestra Remix"
            else:
                query = "T.N.T. AC/DC"

        try:
            import yt_dlp

            ydl_opts = {
                "extract_flat": True,
                "skip_download": True,
                "quiet": True,
                "no_warnings": True,
            }

            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                info = ydl.extract_info(f"ytsearch5:{query}", download=False)

                if "entries" in info and info["entries"]:
                    for entry in info["entries"]:
                        video_id = entry.get("id")
                        title = entry.get("title", query)

                        if video_id:
                            watch_url = f"https://www.youtube.com/watch?v={video_id}&autoplay=1"
                            webbrowser.open(watch_url)
                            return f"Playing {title}."

        except Exception as e:
            print(f"[Media Search Error] Fallback triggered: {e}")

        encoded_query = urllib.parse.quote(query)
        fallback_url = f"https://www.youtube.com/results?search_query={encoded_query}"
        webbrowser.open(fallback_url)
        return f"Searching for {query}."

    def _start_gps_updates(self):
        def updater():
            while True:
                try:
                    g = geocoder.ip("me")
                    if g.ok and g.latlng:
                        self.location = {"lat": g.latlng[0], "lon": g.latlng[1]}
                        self.fetch_weather()
                except Exception as exc:
                    print(f"[GPS Error] {exc}")
                time.sleep(300)

        threading.Thread(target=updater, daemon=True).start()

    def _start_upload_server(self):
        self.upload_app = Flask(__name__)

        @self.upload_app.route("/upload", methods=["POST"])
        def upload_file():
            if "file" not in request.files:
                return jsonify({"error": "no file"}), 400
            f = request.files["file"]
            path = os.path.join("uploads", f.filename)
            os.makedirs(os.path.dirname(path), exist_ok=True)
            f.save(path)
            
            self.load_and_process_file(path)

            return jsonify({"path": path})

        try:
            threading.Thread(
                target=self.upload_app.run,
                kwargs={
                    "host": "127.0.0.1",
                    "port": 5000,
                    "debug": False,
                    "use_reloader": False,
                },
                daemon=True,
            ).start()
        except Exception as e:
            print(f"[Upload Server Error] could not start upload server: {e}")

        @self.upload_app.route("/api/workspace/explore", methods=["GET"])
        def explore_workspace():
            folder_query = request.args.get("folder", "").strip()
            target_dir = Path(folder_query) if os.path.isabs(folder_query) else Path.cwd() / folder_query
            if not target_dir.exists():
                target_dir = Path.cwd()
            
            files = [{"name": item.name, "path": str(item)} for item in target_dir.rglob("*") if item.is_file()]
            return jsonify({"folder": target_dir.name, "files": files})

        @self.upload_app.route("/api/workspace/search-file", methods=["GET"])
        def search_file():
            file_query = request.args.get("file", "").strip()
            search_base = Path.cwd()
            for path in search_base.rglob("*"):
                if path.is_file() and file_query.lower() in path.name.lower():
                    content = path.read_text(encoding="utf-8")
                    return jsonify({"name": path.name, "path": str(path), "content": content, "is_text": True})
            return jsonify({"error": "File not found"}), 404

        @self.upload_app.route("/api/workspace/save-code", methods=["POST"])
        def save_workspace_code():
            data = request.get_json(silent=True) or {}
            file_path = data.get("path", "")
            content = data.get("content", "")

            if not file_path:
                return jsonify({"error": "No path specified"}), 400

            target = Path(file_path)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
            
            if str(target) == self.active_file_path:
                self.update_editor_signal.emit(content)

            return jsonify({"success": True})

    def connect_arms(self):
        try:
            self.robot.connect(self.leader_port.text(), self.follower_port.text())
        except Exception as e:
            print(f"[Hardware Error] {e}")

    def home_arms(self):
        try:
            self.robot.home()
        except Exception as e:
            print(f"[Hardware Error] {e}")

    def set_torque(self, enabled: bool):
        try:
            self.robot.set_torque(enabled)
        except Exception as e:
            print(f"[Hardware Error] {e}")

    def refresh_status(self):
        snapshot = self.robot.snapshot()
        lines = [
            f"{name.title()}: {'ONLINE' if item['connected'] else 'offline'}"
            for name, item in snapshot.items()
        ]
        self.robot_status.setText("\n".join(lines))

    def start_camera(self):
        try:
            if hasattr(self, "cap") and self.cap.isOpened():
                return
            self.cap = cv2.VideoCapture(0)
            if not self.cap.isOpened():
                raise RuntimeError("Failed to open camera device")

            self.cam_timer = QTimer(self)
            self.cam_timer.timeout.connect(self.update_camera_frame)
            self.cam_timer.start(33)
            self.vision_status.setText("Camera live")
        except Exception as e:
            self.vision_status.setText(f"Camera failed ({e})")

    def stop_camera(self):
        try:
            if hasattr(self, "cam_timer") and self.cam_timer.isActive():
                self.cam_timer.stop()
            if hasattr(self, "cap") and self.cap.isOpened():
                self.cap.release()
            self.vision_status.setText("Camera off")
        except Exception as e:
            self.vision_status.setText(f"Camera failed ({e})")

    def update_camera_frame(self):
        if not hasattr(self, "cap") or not self.cap.isOpened():
            return
        ret, frame = self.cap.read()
        if not ret:
            return

        rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        h, w, ch = rgb_frame.shape
        bytes_per_line = ch * w
        qt_image = QImage(rgb_frame.data, w, h, bytes_per_line, QImage.Format.Format_RGB888)
        pixmap = QPixmap.fromImage(qt_image).scaled(
            self.camera_view.size(), Qt.AspectRatioMode.KeepAspectRatio
        )
        self.camera_view.setPixmap(pixmap)

    def closeEvent(self, event):
        self.stop_camera()
        self.robot.disconnect()
        stt = globals().get("stt_thread")
        if stt is not None:
            try:
                stt.stop()
                stt.join(timeout=1)
            except Exception:
                pass
        event.accept()

def fix_known_phrases(text: str) -> str:
    text = text.lower()

    replacements = {
        "hells lightning": "hell's lightning",
        "hals lightning": "hell's lightning",
        "health lightning": "hell's lightning",
        "bearing bit": "bearing bit",
    }

    for wrong, correct in replacements.items():
        text = text.replace(wrong, correct)

    return text

def open_system_app(app_name: str) -> bool:
    import subprocess
    import shutil

    app_name = app_name.lower().strip()

    app_map = {
        "notepad": "notepad.exe",
        "calculator": "calc.exe",
        "calc": "calc.exe",
        "paint": "mspaint.exe",
        "cmd": "cmd.exe",
        "command prompt": "cmd.exe",
        "powershell": "powershell.exe",
        "explorer": "explorer.exe",
    }

    if app_name in app_map:
        app_name = app_map[app_name]

    try:
        subprocess.Popen(app_name)
        return True
    except Exception:
        pass

    found = shutil.which(app_name)
    if found:
        subprocess.Popen(found)
        return True

    return False


app = QApplication(sys.argv)

def main():
    global window

    window = DanWindow()
    window.show()

    global stt_thread
    stt_thread = None
    if whisper is None:
        print("[WHISPER] Whisper is unavailable; speech transcription is disabled.")
    else:
        try:
            stt_thread = STTThread()
            stt_thread.start()
            window.stt_thread = stt_thread
        except Exception as e:
            print(f"[WHISPER] Failed to initialize STT thread: {e}")
            stt_thread = None

    def voice_loop():

        last_valid_command_time = 0
        active_session = False
        session_timeout = 0.0

        while True:
            try:
                spoken_text = stt_queue.get(timeout=0.5)
                if not spoken_text:
                    continue

                spoken_text = spoken_text.strip()
                print(f"[HEARD] {spoken_text}")

                command = extract_wake_word_command(spoken_text)

                if command is not None or (active_session and time.time() < session_timeout):

                    if command == "" and not active_session:
                        greeting = window.generate_wake_greeting()
                        print(f"[DANOVA] {greeting}")

                        with stt_queue.mutex:
                            stt_queue.queue.clear()

                        window.speaker.speak(greeting)

                        active_session = True
                        last_valid_command_time = time.time()

                        while getattr(window.speaker, 'is_speaking', lambda: False)():
                            time.sleep(0.05)

                        session_timeout = time.time() + 6.5
                        continue

                    actual_prompt = command if (command is not None and command != "") else spoken_text
                    active_session = False
                    prompt_lower = actual_prompt.lower()

                    response_text = None

                    if any(prompt_lower.startswith(x) for x in [
                        "play ",
                        "play me ",
                        "find me a video of ",
                        "watch ",
                        "drop my needle"
                    ]):

                        query = (
                            prompt_lower
                            .replace("play me", "")
                            .replace("play", "")
                            .replace("find me a video of", "")
                            .replace("watch", "")
                            .replace("drop my needle", "")
                            .strip()
                        )

                        response_text = window.play_media(query)

                    elif "design" in prompt_lower or "folder" in prompt_lower:

                        try:
                            cleaned = fix_known_phrases(prompt_lower)

                            folder_match = re.search(
                                r"(?:designs?\s+(?:for|in|on)\s+)(.+)",
                                cleaned
                            )

                            file_match = re.search(
                                r"(.+?)\s+(?:from|in|on)\s+(.+)",
                                cleaned
                            )

                            if folder_match:
                                folder_query = re.sub(r"[^\w\s-]", "", folder_match.group(1)).strip()

                                data = window.cloud_client.explore_cloud_folder_dynamic(folder_query)

                                files = data.get("files", [])

                                if files:
                                    window.cloud_files = files

                                    text = f"=== {folder_query.upper()} ===\n\n"
                                    text += "\n".join(
                                        [f"[{i}] {f['name']}" for i, f in enumerate(files)]
                                    )

                                    window.update_editor_signal.emit(text)

                                    response_text = f"Showing {len(files)} files in {folder_query}."
                                else:
                                    response_text = f"No files found in {folder_query}."

                            elif file_match:
                                file_name = re.sub(r"[^\w\s-]", "", file_match.group(1)).strip()
                                folder_query = re.sub(r"[^\w\s-]", "", file_match.group(2)).strip()

                                data = window.cloud_client.explore_cloud_folder_dynamic(folder_query)
                                files = data.get("files", [])

                                match = None
                                for f in files:
                                    if file_name.lower() in f["name"].lower():
                                        match = f
                                        break

                                if match:
                                    urn = window.cloud_client.find_design_urn_globally(match["name"])
                                    if urn:
                                        window.update_viewer_signal.emit(urn)
                                        response_text = f"Opening {match['name']}."
                                    else:
                                        response_text = "Found file but couldn't load model."
                                else:
                                    response_text = f"Couldn't find {file_name} in {folder_query}."

                            else:
                                response_text = "I couldn't understand the design request."

                        except Exception as e:
                            print("[Cloud Error]", e)
                            response_text = "Cloud lookup failed."

                    elif prompt_lower.startswith("open "):

                        target = (
                            prompt_lower
                            .replace("open file", "")
                            .replace("open", "")
                            .replace("the ", "")
                            .strip()
                        )

                        if len(target) >= 2 and open_system_app(target):
                            response_text = f"Opening {target}."
                        else:
                            try:
                                found_file = None
                                for p in Path.cwd().rglob("*"):
                                    if p.is_file() and target in p.name.lower():
                                        found_file = p
                                        break

                                if found_file:
                                    content = analyze_file_content(str(found_file))
                                    window.active_file_path = str(found_file)
                                    window.update_editor_signal.emit(content)
                                    response_text = f"Opened {found_file.name}."
                                else:
                                    response_text = f"Could not find '{target}'."

                            except Exception as e:
                                response_text = f"Error opening file: {e}"

                    if response_text is None:
                        response_text = window.ask_assistant(actual_prompt)

                    print(f"[DANOVA] {response_text}")

                    with stt_queue.mutex:
                        stt_queue.queue.clear()

                    window.speaker.speak(response_text)

                    while window.speaker.is_speaking():
                        time.sleep(0.05)

                    active_session = True
                    session_timeout = time.time() + 6.5

            except queue.Empty:
                if active_session and time.time() > session_timeout:
                    active_session = False
                continue

            except KeyboardInterrupt:
                if stt_thread:
                    stt_thread.stop()
                break

    voice_thread = threading.Thread(target=voice_loop, daemon=True)
    voice_thread.start()
    window.voice_thread = voice_thread

    exit_code = app.exec()

    if stt_thread is not None:
        try:
            stt_thread.stop()
            stt_thread.join(timeout=5)
        except Exception:
            pass

    sys.exit(exit_code)

if __name__ == "__main__":
    main()
